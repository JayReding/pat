from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from export_pdf import HISTORICAL_COLUMNS, OCB_FIELDS, OCB_HEADERS, SUMMARY_LABELS, _MONEY_COLUMNS, _OCB_MONEY

_HEADER_FILL = PatternFill("solid", fgColor="E9ECEF")
_HEADER_FONT = Font(bold=True)
_THIN = Side(style="thin", color="ADB5BD")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_CURRENCY_FORMAT = "$#,##0.00;[Red]($#,##0.00)"
_RANGE_FILL = PatternFill("solid", fgColor="D1E7DD")


def _display_value(value, money=False):
    if value is None:
        return ""
    if isinstance(value, float) and value != value:
        return ""
    if money:
        try:
            return float(value)
        except (TypeError, ValueError):
            return ""
    return value


def build_invoices_workbook(df, sheet_title="Historical Invoices"):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title

    ws.append(list(HISTORICAL_COLUMNS))
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.border = _BORDER

    for _, record in df.iterrows():
        ws.append([
            _display_value(record[c], money=(c in _MONEY_COLUMNS)) for c in HISTORICAL_COLUMNS
        ])

    currency_cols = {
        HISTORICAL_COLUMNS.index(c) + 1 for c in HISTORICAL_COLUMNS if c in _MONEY_COLUMNS
    }

    for row in ws.iter_rows(min_row=2, max_col=len(HISTORICAL_COLUMNS)):
        for cell in row:
            cell.border = _BORDER
            if cell.column in currency_cols and isinstance(cell.value, (int, float)):
                cell.number_format = _CURRENCY_FORMAT

    for i, column in enumerate(HISTORICAL_COLUMNS, start=1):
        letter = get_column_letter(i)
        width = max(
            len(str(column)),
            *(len(str(ws.cell(row=r, column=i).value) or "") for r in range(2, min(ws.max_row, 50) + 1)),
        ) + 2
        ws.column_dimensions[letter].width = min(width, 24)

    ws.freeze_panes = "A2"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_ocb_workbook(df, sheet_title="Ordinary Course (Subjective)",
                       selected_rows=None, summary=None):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]

    summary = summary or {}
    n_sum = len(SUMMARY_LABELS)
    header_row = n_sum + 2
    first_data = header_row + 1

    summary_fill = PatternFill("solid", fgColor="F8F9FA")
    for i, (key, label) in enumerate(SUMMARY_LABELS.items(), start=1):
        label_cell = ws.cell(row=i, column=1, value=label)
        label_cell.font = _HEADER_FONT
        label_cell.fill = summary_fill
        label_cell.border = _BORDER
        value_cell = ws.cell(row=i, column=2)
        value_cell.fill = summary_fill
        value_cell.border = _BORDER
        value = summary.get(key) if summary else None
        if value in (None, "—") or (isinstance(value, float) and value != value):
            value_cell.value = "—"
        elif key == "invoice_count":
            value_cell.value = int(value)
        elif key == "invoice_amount":
            value_cell.value = float(value)
            value_cell.number_format = _CURRENCY_FORMAT
        else:
            value_cell.value = str(value)

    for j, header in enumerate([OCB_HEADERS[f] for f in OCB_FIELDS], start=1):
        cell = ws.cell(row=header_row, column=j, value=header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.border = _BORDER

    for di, (_, record) in enumerate(df.iterrows()):
        r = first_data + di
        for j, field in enumerate(OCB_FIELDS, start=1):
            cell = ws.cell(row=r, column=j, value=_display_value(record[field], money=(field in _OCB_MONEY)))
            cell.border = _BORDER
            if field in _OCB_MONEY and isinstance(cell.value, (int, float)):
                cell.number_format = _CURRENCY_FORMAT

    if selected_rows:
        a, b = sorted((int(selected_rows[0]), int(selected_rows[1])))
        for idx in range(a, min(b, len(df) - 1) + 1):
            for j in range(1, len(OCB_FIELDS) + 1):
                ws.cell(row=first_data + idx, column=j).fill = _RANGE_FILL

    for i, field in enumerate(OCB_FIELDS, start=1):
        letter = get_column_letter(i)
        width = max(
            len(str(OCB_HEADERS[field])),
            *(len(str(ws.cell(row=r, column=i).value) or "") for r in range(first_data, min(ws.max_row, first_data + 49) + 1)),
        ) + 2
        ws.column_dimensions[letter].width = min(width, 24)
    ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width, 26)

    ws.freeze_panes = f"A{first_data}"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
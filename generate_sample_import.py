"""Generate subcase_import_template.xlsx — a fill-in workbook for loading
invoice data into a subcase.

The workbook matches the user-populated columns of the ``invoice_records``
table (see store.py _init_cases_db).  The app derives the four analysis
frames from a single flat invoice table:

    Historical  rows with Payment Date <  petition_date - 90 days
    Preference  rows in [petition_date - 90 days, petition_date]
    Transfers   distinct Transfer Number / Amount / Payment Date in the
                preference window
    New Value   rows with Unpaid = 1

Dates MUST be plain text in YYYY-MM-DD form: the app compares them as
strings when partitioning historical vs preference rows.
"""

from openpyxl import Workbook
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUTPUT = "subcase_import_template.xlsx"

COLUMNS = [
    "Transfer Number",
    "Transfer Amount",
    "Invoice Number",
    "Invoice Amount",
    "Check Amount",
    "Payment Date",
    "Invoice Date",
    "Invoice Due",
    "Terms Days",
    "Days Past Due",
    "WDPD",
    "Invoice to Payment",
    "WI2DEL",
    "Age",
    "Unpaid",
    "Check Date",
]

REQUIRED = {
    "Transfer Number",
    "Transfer Amount",
    "Invoice Number",
    "Invoice Amount",
    "Payment Date",
    "Invoice Date",
}

MONEY_COLUMNS = {"Transfer Amount", "Invoice Amount", "Check Amount"}
INTEGER_COLUMNS = {"Terms Days", "Days Past Due", "Invoice to Payment", "Unpaid"}

_HEADER_REQUIRED_FILL = PatternFill("solid", fgColor="DDEBF7")  # light blue
_HEADER_OPTIONAL_FILL = PatternFill("solid", fgColor="F2F2F2")  # light gray
_HEADER_FONT = Font(bold=True)
_THIN = Side(style="thin", color="ADB5BD")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_CURRENCY_FORMAT = "$#,##0.00;[Red]($#,##0.00)"

# Hypothetical petition date 2023-10-15 => preference window 2023-07-17..2023-10-15.
# Rows 1-3 land in the historical window; rows 4-9 in the preference window;
# rows 8-9 are flagged Unpaid=1 and feed the New Value frame.
EXAMPLE_ROWS = [
    # (Transfer No, Transfer Amt, Invoice No, Invoice Amt, Check Amt, Payment Date,
    #  Invoice Date, Invoice Due, Terms Days, Days Past Due, WDPD, Invoice to Payment,
    #  WI2DEL, Age, Unpaid, Check Date)
    ("T-0001", 750.00,  "INV-1001", 750.00,  750.00,  "2023-05-10", "2023-04-25", "2023-05-25", 30, 0, 0, 15, 15, 15, 0, "2023-05-10"),
    ("T-0002", 1250.00, "INV-1002", 1250.00, 1250.00, "2023-06-06", "2023-05-10", "2023-06-09", 30, 0, 0, 27, 27, 27, 0, "2023-06-06"),
    ("T-0003", 500.00,  "INV-1003", 500.00,  500.00,  "2023-06-28", "2023-06-01", "2023-07-01", 30, 0, 0, 27, 27, 27, 0, "2023-06-28"),
    ("T-0004", 900.00,  "INV-1004", 900.00,  900.00,  "2023-07-20", "2023-07-02", "2023-08-01", 30, 0, 0, 18, 18, 18, 0, "2023-07-20"),
    ("T-0005", 1500.00, "INV-1005", 1500.00, 1500.00, "2023-08-12", "2023-07-19", "2023-08-18", 30, 0, 0, 24, 24, 24, 0, "2023-08-12"),
    ("T-0006", 800.00,  "INV-1006", 800.00,  800.00,  "2023-09-03", "2023-08-15", "2023-09-14", 30, 0, 0, 19, 19, 19, 0, "2023-09-03"),
    ("T-0007", 1100.00, "INV-1007", 1100.00, 1100.00, "2023-09-20", "2023-09-01", "2023-10-01", 30, 0, 0, 19, 19, 19, 0, "2023-09-20"),
    ("T-0008", 650.00,  "INV-1008", 650.00,  None,    "2023-09-25", "2023-09-02", "2023-10-02", 30, 0, 0, 23, 23, 23, 1, "2023-09-25"),
    ("T-0009", 420.00,  "INV-1009", 420.00,  None,    "2023-10-06", "2023-09-15", "2023-10-15", 30, 0, 0, 21, 21, 21, 1, "2023-10-06"),
]

INSTRUCTIONS = [
    "SUBCASE DATA IMPORT TEMPLATE",
    "",
    "About this file",
    "  This workbook holds the invoice records for ONE subcase.  The subcase it",
    "  belongs to is chosen in the app at import time; this file only carries",
    "  the invoice rows.",
    "",
    "Getting started",
    "  1. Delete the 9 example rows before filling in your own data.",
    "  2. Keep the header row exactly as it is (row 1) - the importer matches",
    "     columns by header name.",
    "  3. Dates MUST be entered as plain text in YYYY-MM-DD form",
    "     (e.g. 2023-10-15).  A consistent text format is required because the",
    "     app compares dates as strings when splitting periods.",
    "",
    "Columns",
    "  Blue headers are required; gray headers are optional.",
    "  Required:",
    "    Transfer Number / Transfer Amount / Invoice Number / Invoice Amount /",
    "    Payment Date / Invoice Date",
    "  Optional but strongly recommended:",
    "    Invoice to Payment - days from invoice to payment; the weighted DSO",
    "      metric used in the analysis (payment date minus invoice date).",
    "    Days Past Due - the weighted DPD metric used in the analysis.",
    "  Optional:",
    "    Check Amount, Invoice Due, Terms Days, WDPD, WI2DEL, Age, Check Date",
    "    Unpaid - enter 1 to include the invoice in the New Value frame;",
    "             leave blank (or 0) otherwise. Defaults to 0.",
    "",
    "How rows are partitioned",
    "  The preference window is the 90 days ending on the main case's petition",
    "  date (which comes from the linked main case, not from this file):",
    "    Historical : Payment Date <  petition date - 90 days",
    "    Preference : petition date - 90 days  <=  Payment Date  <=  petition date",
    "    New Value  : Unpaid = 1",
    "  The example rows use a hypothetical petition date of 2023-10-15, so rows",
    "  1-3 are historical and rows 4-9 fall in the preference window.",
    "",
    "Tips",
    "  Give every invoice its Transfer Number so invoices that were paid by the",
    "  same transfer group together in the Transfer tab.",
    "  Amounts are read as currency; enter plain numbers (no $ or commas).",
]

MONEY_INDEXES = {COLUMNS.index(c) + 1 for c in MONEY_COLUMNS}
INTEGER_INDEXES = {COLUMNS.index(c) + 1 for c in INTEGER_COLUMNS}


def _write_invoices(ws):
    ws.append(COLUMNS)
    for i, header in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=i, value=header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_REQUIRED_FILL if header in REQUIRED else _HEADER_OPTIONAL_FILL
        cell.border = _BORDER

    for row in EXAMPLE_ROWS:
        ws.append(row)

    for row in ws.iter_rows(min_row=2, max_col=len(COLUMNS)):
        for cell in row:
            cell.border = _BORDER
            if cell.column in MONEY_INDEXES and isinstance(cell.value, (int, float)):
                cell.number_format = _CURRENCY_FORMAT

    dv = DataValidation(
        type="list", formula1='"0,1"', allow_blank=True,
        showErrorMessage=True, errorTitle="Invalid Unpaid value",
        error="Enter 0 (or leave blank) for paid invoices, or 1 to include the invoice in New Value.",
    )
    unpaid_col = COLUMNS.index("Unpaid") + 1
    ws.add_data_validation(dv)
    dv.add(f"{get_column_letter(unpaid_col)}2:{get_column_letter(unpaid_col)}{ws.max_row}")

    for i, column in enumerate(COLUMNS, start=1):
        letter = get_column_letter(i)
        width = max(
            len(str(column)),
            *(len(str(ws.cell(row=r, column=i).value) or "")
              for r in range(2, ws.max_row + 1)),
        ) + 2
        ws.column_dimensions[letter].width = min(width, 26)

    ws.freeze_panes = "A2"


def _write_instructions(ws):
    ws.column_dimensions["A"].width = 110
    for i, line in enumerate(INSTRUCTIONS, start=1):
        cell = ws.cell(row=i, column=1, value=line if line else None)
        if line and not line.startswith("  ") and line != "":
            cell.font = Font(bold=True, size=11)


def build_sample_workbook():
    wb = Workbook()
    invoices = wb.active
    assert invoices is not None
    invoices.title = "Invoices"
    _write_invoices(invoices)
    _write_instructions(wb.create_sheet("Instructions"))
    return wb


def main():
    wb = build_sample_workbook()
    wb.save(OUTPUT)
    print(f"Wrote {OUTPUT}")

    invoices = wb["Invoices"]
    headers = [c.value for c in invoices[1]]
    print(f"Sheets: {wb.sheetnames}")
    print(f"Header row: {headers}")
    print(f"Data rows (incl. example): {invoices.max_row - 1}")
    print(f"Required columns: {sorted(REQUIRED)}")


if __name__ == "__main__":
    main()